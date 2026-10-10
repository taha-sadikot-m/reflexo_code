"""Generate the HolonCode document pack used to test voice cross-channel media.

Each asset stays small on purpose: extraction is capped at 30s in
shared/document_processor.py and the voice tool clips extracted text at 1500
characters, so the values worth quizzing the agent on sit near the top.
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

OUT_DIR = ROOT / "tests" / "assets" / "voice_doc_pack"

# fpdf2 core fonts are latin-1 only, so every asset stays ASCII (INR, not the
# rupee sign). Docling and the OCR fallbacks both handle that fine.
PDF_LINES = [
    ("h1", "HolonCode - Services and Pricing 2026"),
    ("b", "Support email: support@holoncode.dev"),
    ("b", "Support phone: +91 80 4718 2290"),
    ("b", "Price list code: HC-PRICE-2026-04"),
    ("gap", ""),
    ("h2", "Services"),
    ("p", "Starter Website Build - INR 45,000 - delivery in 2 weeks"),
    ("p", "WhatsApp AI Agent Setup - INR 78,500 - delivery in 3 weeks"),
    ("p", "Voice Agent Add-on - INR 32,000 - delivery in 1 week"),
    ("p", "Annual Support Retainer - INR 12,000 per month"),
    ("gap", ""),
    ("h2", "Terms"),
    ("p", "All prices exclude 18 percent GST."),
    ("p", "Quotes stay valid for 30 days from the issue date."),
    ("p", "Onboarding starts within 3 business days of the advance payment."),
]

DOCX_BLOCKS = [
    ("h1", "HolonCode - Client Onboarding Note"),
    ("p", "Document code: HC-ONBOARD-118"),
    ("h2", "Business hours"),
    ("p", "Monday to Friday: 9:30 AM to 6:30 PM IST"),
    ("p", "Saturday: 10:00 AM to 2:00 PM IST"),
    ("p", "Sunday: closed"),
    ("h2", "Refund policy"),
    (
        "p",
        "Advance payments are refundable within 7 days of project kickoff, "
        "minus work already delivered. After 7 days the advance is adjusted "
        "against the final invoice.",
    ),
    ("h2", "Escalation"),
    ("p", "Escalation contact: priya.nair@holoncode.dev"),
]

XLSX_HEADER = ["SKU", "Item", "Unit Price INR", "Lead Time Days"]
XLSX_ROWS = [
    ["HC-101", "Landing Page Design", 18000, 5],
    ["HC-102", "WhatsApp Agent Setup", 78500, 21],
    ["HC-103", "Voice Agent Add-on", 32000, 7],
    ["HC-104", "CRM Integration", 26500, 10],
    ["HC-105", "Analytics Dashboard", 21000, 8],
    ["HC-106", "Support Retainer", 12000, 1],
]

CARD_TITLE = "HolonCode Price Card"
CARD_CODE = "Card code: HC-CARD-77"
CARD_ITEMS = [
    "Landing Page Design      INR 18,000",
    "WhatsApp Agent Setup     INR 78,500",
    "Voice Agent Add-on       INR 32,000",
    "Support Retainer         INR 12,000",
]
CARD_FOOTER = "Offer valid till 30 September 2026"


def _write_pdf(path: Path) -> None:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    styles = {
        "h1": ("B", 16, 9),
        "h2": ("B", 13, 8),
        "b": ("B", 11, 7),
        "p": ("", 11, 7),
    }
    pdf = FPDF(format="A4", unit="mm")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    for kind, text in PDF_LINES:
        if kind == "gap":
            pdf.ln(4)
            continue
        style, size, height = styles[kind]
        pdf.set_font("Helvetica", style, size)
        # Default multi_cell leaves the cursor at the right margin, which
        # starves the next full-width cell.
        pdf.multi_cell(0, height, text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        if kind == "h1":
            pdf.ln(2)
    pdf.output(str(path))


def _write_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    for kind, text in DOCX_BLOCKS:
        if kind == "h1":
            doc.add_heading(text, level=1)
        elif kind == "h2":
            doc.add_heading(text, level=2)
        else:
            doc.add_paragraph(text)
    doc.save(str(path))


def _write_xlsx(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    sheet = book.active
    sheet.title = "Price List"
    sheet.append(XLSX_HEADER)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in XLSX_ROWS:
        sheet.append(row)
    for column, width in zip("ABCD", (12, 26, 16, 16)):
        sheet.column_dimensions[column].width = width
    book.save(str(path))


def _card_font(size: int):
    from PIL import ImageFont

    for name in ("arialbd.ttf", "arial.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    try:
        return ImageFont.load_default(size=size)
    except TypeError:
        return ImageFont.load_default()


def _write_png(path: Path) -> None:
    from PIL import Image, ImageDraw

    width, height = 1100, 600
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle([(20, 20), (width - 20, height - 20)], outline="black", width=4)

    title_font = _card_font(52)
    body_font = _card_font(38)
    small_font = _card_font(30)

    y = 70
    draw.text((60, y), CARD_TITLE, fill="black", font=title_font)
    y += 80
    draw.text((60, y), CARD_CODE, fill="black", font=small_font)
    y += 70
    draw.line([(60, y), (width - 60, y)], fill="black", width=3)
    y += 30
    for item in CARD_ITEMS:
        draw.text((60, y), item, fill="black", font=body_font)
        y += 60
    y += 20
    draw.text((60, y), CARD_FOOTER, fill="black", font=small_font)
    image.save(str(path), format="PNG")


BUILDERS = [
    ("HolonCode_Services.pdf", _write_pdf),
    ("HolonCode_Onboarding_Note.docx", _write_docx),
    ("HolonCode_Price_List.xlsx", _write_xlsx),
    ("HolonCode_Price_Card.png", _write_png),
]


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for filename, builder in BUILDERS:
        path = OUT_DIR / filename
        builder(path)
        print(f"wrote {path.relative_to(ROOT)} ({path.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
